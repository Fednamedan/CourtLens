"""Экспорт результатов: CSV строго по образцу + XLSX с подсветкой и листом замечаний."""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import OCR_COLUMNS, ExtractionResult, FieldStatus
from app.validate.checks import issues

RED = PatternFill("solid", fgColor="F8CBAD")      # не извлечено
YELLOW = PatternFill("solid", fgColor="FFE699")   # низкая уверенность / не прошло проверку
ORANGE = PatternFill("solid", fgColor="F4B084")   # весь документ требует юриста

DATA_COLUMNS = [c for c in OCR_COLUMNS if c not in ("тип документа", "Файл")]


def result_to_row(res: ExtractionResult) -> dict[str, object]:
    row: dict[str, object] = {"тип документа": res.doc_type.value if res.doc_type.value != "unknown" else "",
                              "Файл": res.file}
    for col in DATA_COLUMNS:
        fv = res.fields.get(col)
        # INVALID и пустые значения не попадают в таблицу как «достоверные» - только пометка в замечаниях
        row[col] = fv.value if fv and fv.status in (FieldStatus.OK, FieldStatus.LOW_CONFIDENCE) else ""
    return row


def write_csv(results: list[ExtractionResult], path: str) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=OCR_COLUMNS)
        w.writeheader()
        for r in results:
            w.writerow(result_to_row(r))


def write_xlsx(results: list[ExtractionResult], path: str) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Результат"
    headers = OCR_COLUMNS + ["юрист", "замечания"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for res in results:
        row = result_to_row(res)
        probs = issues(res, DATA_COLUMNS)
        ws.append([row[c] for c in OCR_COLUMNS]
                  + ["да" if res.needs_lawyer else "", "; ".join(probs)])
        r = ws.max_row
        for i, col in enumerate(OCR_COLUMNS, start=1):
            fv = res.fields.get(col)
            if col in DATA_COLUMNS:
                if fv is None or fv.status == FieldStatus.MISSING:
                    ws.cell(r, i).fill = RED
                elif fv.status != FieldStatus.OK:
                    ws.cell(r, i).fill = YELLOW
        if res.needs_lawyer:
            ws.cell(r, len(OCR_COLUMNS) + 1).fill = ORANGE
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(h) + 2, 12), 40)
    ws.freeze_panes = "C2"

    det = wb.create_sheet("Детали")
    det.append(["Файл", "поле", "значение", "статус", "уверенность", "стр.", "цитата", "примечание"])
    for res in results:
        for name, fv in res.fields.items():
            det.append([res.file, name, fv.value, fv.status.value, fv.confidence, fv.page,
                        fv.quote, fv.note])
    lg = wb.create_sheet("Для юриста")
    lg.append(["Файл", "фраза-маркер", "стр.", "сходство", "контекст"])
    for res in results:
        for fl in res.legal_flags:
            lg.append([res.file, fl.phrase, fl.page, round(fl.score, 2), fl.context])
    wb.save(path)


# ---------------- таблица ФССП (XML) ----------------
from app.models import XML_COLUMNS  # noqa: E402


def xml_row(res: ExtractionResult) -> dict[str, object]:
    row: dict[str, object] = {"FileName": res.file}
    for col in XML_COLUMNS[1:]:
        fv = res.fields.get(col)
        row[col] = fv.value if fv and fv.status in (FieldStatus.OK, FieldStatus.LOW_CONFIDENCE) else ""
    return row


def write_xml_csv(results: list[ExtractionResult], path: str) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=XML_COLUMNS)
        w.writeheader()
        for r in results:
            w.writerow(xml_row(r))


def write_xml_xlsx(results: list[ExtractionResult], path: str) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Результат"
    headers = XML_COLUMNS + ["замечания"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for res in results:
        row = xml_row(res)
        notes = [f"{n}: {fv.note}" for n, fv in res.fields.items()
                 if fv.note and (fv.status != FieldStatus.OK or n.startswith(("date_", "rub_")))
                 and fv.status != FieldStatus.OK]
        ws.append([row[c] for c in XML_COLUMNS] + ["; ".join(notes)])
        r = ws.max_row
        for i, col in enumerate(XML_COLUMNS[1:], start=2):
            fv = res.fields.get(col)
            if col in ("rub_post", "rub_poshlina", "rub_peni"):
                continue  # пустая необязательная сумма - норма, не ошибка
            if fv is None or fv.status == FieldStatus.MISSING:
                ws.cell(r, i).fill = RED
            elif fv.status != FieldStatus.OK:
                ws.cell(r, i).fill = YELLOW
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(h) + 2, 12), 40)
    ws.freeze_panes = "B2"
    det = wb.create_sheet("Детали")
    det.append(["Файл", "поле", "значение", "статус", "уверенность", "источник", "примечание"])
    for res in results:
        for name, fv in res.fields.items():
            det.append([res.file, name, fv.value, fv.status.value, fv.confidence, fv.quote, fv.note])
    wb.save(path)


def write_debtors_csv(results: list[ExtractionResult], path: str) -> int:
    """Таблица «по должникам»: для актов с блоками должников - по строке на каждого. Возвращает число строк."""
    rows = []
    for r in results:
        for i, fields in enumerate(r.debtor_rows, start=1):
            tmp = ExtractionResult(file=r.file, doc_type=r.doc_type, fields=fields)
            row = result_to_row(tmp)
            row["должник №"] = i
            rows.append(row)
    if not rows:
        return 0
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=OCR_COLUMNS + ["должник №"])
        w.writeheader()
        w.writerows(rows)
    return len(rows)
